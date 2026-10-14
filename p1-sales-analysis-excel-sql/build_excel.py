from pathlib import Path
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs"
DEST = ROOT / "Sales_Analysis.xlsx"


def add_frame(wb: Workbook, title: str, frame: pd.DataFrame) -> None:
    ws = wb.create_sheet(title=title[:31])
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        ws.append([None if pd.isna(v) else v.item() if hasattr(v, "item") else v for v in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="183B56")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = max((len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200]), default=10)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 34)


def main() -> None:
    if not (OUT / "kpi_overview.csv").is_file():
        sys.exit("Analysis outputs not found. Run `python analysis.py` first.")
    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Summary"
    ws.append(["Olist E-commerce Sales Analysis", "Generated from the public source data"])
    ws.append(["Metric", "Value"])
    kpi = pd.read_csv(OUT / "kpi_overview.csv")
    if not kpi.empty:
        for col in kpi.columns:
            ws.append([col.replace("_", " ").title(), kpi.iloc[0][col]])
    ws.append([])
    ws.append(["Interpretation boundary", "Product sales and freight are separate; neither is profit."])
    ws.append(["Dataset", "Olist Brazilian E-Commerce Public Dataset"])
    ws.append(["Source", "https://www.kaggle.com/datasets/olistbr/brazilian-ecommerce"])
    for cell in ws[1]:
        cell.font = Font(bold=True, size=14, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="183B56")
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="287271")
    ws.column_dimensions["A"].width = 43
    ws.column_dimensions["B"].width = 80
    ws.freeze_panes = "A3"

    tables = [
        ("Monthly Sales", "monthly_sales.csv"),
        ("Categories", "category_performance.csv"),
        ("YoY Growth", "category_yoy_growth.csv"),
        ("Pareto", "category_pareto.csv"),
        ("Customer RFM", "customer_rfm.csv"),
        ("Delivery", "delivery_performance.csv"),
        ("Payment Reconciliation", "payment_reconciliation.csv"),
        ("Data Quality", "data_quality_report.csv"),
    ]
    for title, filename in tables:
        path = OUT / filename
        if path.exists():
            add_frame(wb, title, pd.read_csv(path))

    if "Monthly Sales" in wb.sheetnames:
        src = wb["Monthly Sales"]
        chart = LineChart()
        chart.title = "Monthly Product Sales"
        chart.y_axis.title = "BRL"
        chart.x_axis.title = "Purchase Month"
        chart.height, chart.width = 8, 16
        chart.add_data(Reference(src, min_col=3, min_row=1, max_row=src.max_row), titles_from_data=True)
        chart.set_categories(Reference(src, min_col=1, min_row=2, max_row=src.max_row))
        wb["Executive Summary"].add_chart(chart, "D2")
    wb.save(DEST)
    print(f"Workbook written to {DEST}")


if __name__ == "__main__":
    main()
